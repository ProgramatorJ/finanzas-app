import os
import json
import datetime
import math
import openpyxl
from google.oauth2.credentials import Credentials
from google.cloud import firestore

# --- Configuración de Autenticación de Firebase ---
print("Cargando credenciales de Firebase CLI...")
config_path = os.path.expanduser(r"~\.config\configstore\firebase-tools.json")
if not os.path.exists(config_path):
    raise FileNotFoundError(f"No se encontró la configuración del Firebase CLI en {config_path}. ¿Ejecutaste 'firebase login'?")

with open(config_path, "r", encoding="utf-8") as f:
    config = json.load(f)

tokens = config.get("tokens", {})
access_token = tokens.get("access_token")
refresh_token = tokens.get("refresh_token")

client_id = "563584335869-uu3v1rfjqc2msthab9xgbe2nt9gg7ddp.apps.googleusercontent.com"
client_secret = "j9Te1-i5g3gJnS5s1RLk98Ls"

creds = Credentials(
    token=access_token,
    refresh_token=refresh_token,
    token_uri="https://oauth2.googleapis.com/token",
    client_id=client_id,
    client_secret=client_secret
)

db = firestore.Client(project="mi-gestor-prestamos", credentials=creds)
print("Conectado con éxito a Firestore!")

# Obtener un ID de usuario de administrador de la colección 'users' para asociar los registros creados
print("Buscando usuario administrador en la base de datos...")
users = list(db.collection("users").limit(1).stream())
if users:
    user_uid = users[0].id
    print(f"Usuario asignado para 'createdByUid': {user_uid} ({users[0].to_dict().get('email')})")
else:
    user_uid = "migration_admin"
    print("No se encontraron usuarios, usando ID por defecto: 'migration_admin'")

# --- Funciones de Normalización ---
def clean_str(val):
    if val is None or str(val).strip().lower() in ["none", "null", "n/a", "#n/a"]:
        return ""
    return str(val).strip()

def normalize_id(val):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str.lower() in ["none", "null", "n/a", "#n/a"]:
        return ""
    if "." in val_str:
        try:
            return str(int(float(val_str)))
        except ValueError:
            pass
    return val_str

def normalize_name(val):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str.lower() in ["none", "null", "n/a", "#n/a"]:
        return ""
    return val_str

def clean_phone(val):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str.lower() in ["none", "null", "n/a", "#n/a", "0", "0.0"]:
        return ""
    if "." in val_str:
        try:
            return str(int(float(val_str)))
        except ValueError:
            pass
    return val_str

# --- Funciones de Calendario Comercial 360 ---
def last_day_of_month(year, month):
    if month == 12:
        return 31
    return (datetime.datetime(year, month + 1, 1) - datetime.timedelta(days=1)).day

def add_commercial_days(dt, days):
    day = dt.day
    month = dt.month
    year = dt.year

    if day > 30:
        day = 30

    total_days = (year * 360) + ((month - 1) * 30) + day + days

    new_year = total_days // 360
    remainder = total_days % 360

    new_month = (remainder // 30) + 1
    new_day = remainder % 30

    if new_day == 0:
        new_day = 30
        new_month -= 1
        if new_month == 0:
            new_month = 12
            new_year -= 1

    real_last_day = last_day_of_month(new_year, new_month)
    if new_day > real_last_day:
        new_day = real_last_day

    # Conservar el mismo componente de hora
    return datetime.datetime(new_year, new_month, new_day, dt.hour, dt.minute, dt.second)

def generate_installment_dates(first_installment_date, number_of_installments, frequency):
    if frequency == 'weekly':
        interval_days = 7
    elif frequency == 'biweekly':
        interval_days = 15
    else:
        interval_days = 30

    dates = []
    current = first_installment_date
    for i in range(number_of_installments):
        if i == 0:
            dates.append(current)
        else:
            current = add_commercial_days(current, interval_days)
            dates.append(current)
    return dates

def get_target_installment_number(cuota_val):
    if cuota_val is None:
        return 1
    cuota_str = str(cuota_val).strip()
    if '-' in cuota_str:
        parts = cuota_str.split('-')
        try:
            return int(float(parts[0]))
        except ValueError:
            pass
    try:
        return int(float(cuota_str))
    except ValueError:
        return 1

# --- Cargar Excel ---
excel_path = r"C:\Users\Windows 10\Desktop\projects Antigravity\Copia de BANX.xlsx"
print(f"Cargando archivo Excel desde {excel_path}...")
wb = openpyxl.load_workbook(excel_path, data_only=True)

# Colores a omitir
PINK_HEX = "FFC27BA0"
GRAY_HEX = "FF434343"

# --- 1. Leer y Filtrar los Créditos válidos de la pestaña 'Credito' ---
print("\n--- Analizando créditos en la pestaña 'Credito' ---")
ws_credito = wb['Credito']

valid_credits = []
valid_client_ids = set()

for r_idx in range(2, ws_credito.max_row + 1):
    c_code = ws_credito.cell(row=r_idx, column=1)
    c_name = ws_credito.cell(row=r_idx, column=4)
    c_cedula = ws_credito.cell(row=r_idx, column=5)
    c_valor = ws_credito.cell(row=r_idx, column=6)
    c_plazo = ws_credito.cell(row=r_idx, column=7)
    c_periodo = ws_credito.cell(row=r_idx, column=9)
    c_fecha = ws_credito.cell(row=r_idx, column=2)
    
    val_name = c_name.value
    val_code = c_code.value
    val_cedula = c_cedula.value
    
    if val_name is None or val_code is None or val_cedula is None:
        continue
        
    # Verificar color
    font_color = c_name.font.color.rgb if c_name.font and c_name.font.color else None
    
    if font_color == PINK_HEX:
        print(f"Omitiendo fila {r_idx:2d} (ROSADO): Code={val_code}, Name={val_name}")
        continue
    if font_color == GRAY_HEX:
        print(f"Omitiendo fila {r_idx:2d} (GRIS): Code={val_code}, Name={val_name}")
        continue
        
    code_str = normalize_id(val_code)
    client_id = normalize_id(val_cedula)
    principal = float(c_valor.value) if c_valor.value is not None else 0.0
    plazo = int(float(c_plazo.value)) if c_plazo.value is not None else 1
    periodo_val = float(c_periodo.value) if c_periodo.value is not None else 1.0
    
    # Frecuencia
    if periodo_val == 1.0:
        freq = 'monthly'
    elif periodo_val == 2.0:
        freq = 'biweekly'
    elif periodo_val == 4.0:
        freq = 'weekly'
    else:
        freq = 'monthly'
        
    # Fecha desembolso
    fecha_val = c_fecha.value
    if isinstance(fecha_val, datetime.datetime):
        disb_date = fecha_val
    elif isinstance(fecha_val, str):
        try:
            disb_date = datetime.datetime.strptime(fecha_val, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            disb_date = datetime.datetime.now()
    else:
        disb_date = datetime.datetime.now()
        
    valid_credits.append({
        'creditId': code_str,
        'clientId': client_id,
        'clientName': normalize_name(val_name),
        'principalAmount': principal,
        'termInMonths': plazo,
        'paymentFrequency': freq,
        'disbursementDate': disb_date
    })
    valid_client_ids.add(client_id)
    print(f"Crédito Válido detectado en fila {r_idx:2d} ({font_color}): Code={code_str}, Name={val_name}, Cédula={client_id}")

print(f"\nTotal créditos válidos: {len(valid_credits)}")
print(f"Total cédulas únicas asociadas: {len(valid_client_ids)}")

# --- 2. Leer y Filtrar Clientes de la pestaña 'Clientes' ---
print("\n--- Analizando clientes en la pestaña 'Clientes' ---")
ws_clientes = wb['Clientes']

clients_to_import = []
found_client_ids = set()

for r_idx in range(2, ws_clientes.max_row + 1):
    c_cedula = ws_clientes.cell(row=r_idx, column=1)
    c_nombre = ws_clientes.cell(row=r_idx, column=2)
    c_telefono = ws_clientes.cell(row=r_idx, column=3)
    c_direccion = ws_clientes.cell(row=r_idx, column=4)
    c_comentario = ws_clientes.cell(row=r_idx, column=6)
    
    val_cedula = normalize_id(c_cedula.value)
    val_nombre = normalize_name(c_nombre.value)
    
    if not val_cedula or not val_nombre:
        continue
        
    # Solo importar si está en la lista de clientes con créditos activos/válidos
    if val_cedula in valid_client_ids:
        clients_to_import.append({
            'clientId': val_cedula,
            'fullName': val_nombre,
            'idNumber': val_cedula,
            'phone': clean_phone(c_telefono.value),
            'address': clean_str(c_direccion.value) if c_direccion.value else "",
            'assignedCollectorIds': [],
            'createdByUid': user_uid,
            'isActive': True,
            'notes': clean_str(c_comentario.value) if c_comentario.value else None
        })
        found_client_ids.add(val_cedula)
        print(f"Cliente Válido detectado en fila {r_idx:3d}: ID={val_cedula}, Nombre={val_nombre}")

# Agregar clientes que tienen crédito pero no estaban en la lista de Clientes (resiliencia)
missing_clients = valid_client_ids - found_client_ids
for m_id in missing_clients:
    # Buscar nombre del crédito correspondiente
    name = "Desconocido"
    for vc in valid_credits:
        if vc['clientId'] == m_id:
            name = vc['clientName']
            break
    clients_to_import.append({
        'clientId': m_id,
        'fullName': name,
        'idNumber': m_id,
        'phone': "",
        'address': "",
        'assignedCollectorIds': [],
        'createdByUid': user_uid,
        'isActive': True,
        'notes': "Creado automáticamente por migración (crédito existente sin ficha)"
    })
    print(f"Cliente Faltante agregado automáticamente: ID={m_id}, Nombre={name}")

print(f"Total clientes a importar: {len(clients_to_import)}")



# --- 3. Leer Pagos de la pestaña 'Pago Cuota' ---
print("\n--- Analizando pagos en la pestaña 'Pago Cuota' ---")
ws_pago = wb['Pago Cuota']

all_payments = []
for r_idx in range(2, ws_pago.max_row + 1):
    c_prog_fecha = ws_pago.cell(row=r_idx, column=1) # Column A
    c_fecha2 = ws_pago.cell(row=r_idx, column=2)     # Column B
    c_codigo = ws_pago.cell(row=r_idx, column=3)     # Column C
    c_nombre = ws_pago.cell(row=r_idx, column=4)     # Column D
    c_valor = ws_pago.cell(row=r_idx, column=5)      # Column E
    c_cuota = ws_pago.cell(row=r_idx, column=9)      # Column I
    
    val_codigo = normalize_id(c_codigo.value)
    if not val_codigo:
        continue
        
    # Verificar si pertenece a uno de nuestros créditos kept
    is_matching = False
    for vc in valid_credits:
        if vc['creditId'] == val_codigo:
            is_matching = True
            break
            
    if is_matching:
        valor = float(c_valor.value) if c_valor.value is not None else 0.0
        cuota_num = get_target_installment_number(c_cuota.value)
        
        fecha_val = c_fecha2.value
        if isinstance(fecha_val, datetime.datetime):
            fecha2 = fecha_val
        elif isinstance(fecha_val, str):
            try:
                fecha2 = datetime.datetime.strptime(fecha_val, "%Y-%m-%d %H:%M:%S")
            except ValueError:
                fecha2 = datetime.datetime.now()
        else:
            fecha2 = datetime.datetime.now()
            
        all_payments.append({
            'creditId': val_codigo,
            'fecha_2': fecha2,
            'valor': valor,
            'cuota_num': cuota_num,
            'nombre': normalize_name(c_nombre.value)
        })

print(f"Total pagos coincidentes en la hoja: {len(all_payments)}")

# --- EJECUCIÓN DE LA MIGRACIÓN ---
print("\n==================================================")
print("INICIANDO PROCESO DE ESCRITURA EN FIRESTORE")
print("==================================================")

# 1. Limpieza de Firestore
print("Limpiando colecciones anteriores de Firestore...")

# Borrar todas las cuotas (installments)
print("  Eliminando subcolección group 'installments'...")
inst_docs = list(db.collection_group("installments").stream())
if inst_docs:
    batch = db.batch()
    for idx, doc in enumerate(inst_docs):
        batch.delete(doc.reference)
        if (idx + 1) % 400 == 0:
            batch.commit()
            batch = db.batch()
    batch.commit()
print(f"  [OK] {len(inst_docs)} cuotas eliminadas.")

# Borrar todos los pagos (payments)
print("  Eliminando subcolección group 'payments'...")
pay_docs = list(db.collection_group("payments").stream())
if pay_docs:
    batch = db.batch()
    for idx, doc in enumerate(pay_docs):
        batch.delete(doc.reference)
        if (idx + 1) % 400 == 0:
            batch.commit()
            batch = db.batch()
    batch.commit()
print(f"  [OK] {len(pay_docs)} pagos eliminados.")

# Borrar todos los créditos (credits)
print("  Eliminando colección 'credits'...")
credit_docs = list(db.collection("credits").stream())
if credit_docs:
    batch = db.batch()
    for idx, doc in enumerate(credit_docs):
        batch.delete(doc.reference)
        if (idx + 1) % 400 == 0:
            batch.commit()
            batch = db.batch()
    batch.commit()
print(f"  [OK] {len(credit_docs)} créditos eliminados.")

# Borrar todos los clientes (clients)
print("  Eliminando colección 'clients'...")
client_docs = list(db.collection("clients").stream())
if client_docs:
    batch = db.batch()
    for idx, doc in enumerate(client_docs):
        batch.delete(doc.reference)
        if (idx + 1) % 400 == 0:
            batch.commit()
            batch = db.batch()
    batch.commit()
print(f"  [OK] {len(client_docs)} clientes eliminados.")

# Borrar citas
print("  Eliminando colección 'appointments'...")
appt_docs = list(db.collection("appointments").stream())
if appt_docs:
    batch = db.batch()
    for idx, doc in enumerate(appt_docs):
        batch.delete(doc.reference)
        if (idx + 1) % 400 == 0:
            batch.commit()
            batch = db.batch()
    batch.commit()
print(f"  [OK] {len(appt_docs)} citas eliminadas.")

# 2. Guardar Clientes
print("\nGuardando nuevos clientes...")
batch = db.batch()
for idx, cl in enumerate(clients_to_import):
    cl_ref = db.collection("clients").document(cl['clientId'])
    cl_data = cl.copy()
    cl_data['createdAt'] = datetime.datetime.now()
    cl_data['updatedAt'] = datetime.datetime.now()
    batch.set(cl_ref, cl_data)
    if (idx + 1) % 400 == 0:
        batch.commit()
        batch = db.batch()
batch.commit()
print(f"[OK] {len(clients_to_import)} clientes importados correctamente.")

# 3. Guardar Créditos, Cuotas y aplicar Pagos
print("\nGuardando créditos, plan de cuotas y aplicando pagos...")

for cred in valid_credits:
    c_id = cred['creditId']
    cl_id = cred['clientId']
    c_name = cred['clientName']
    principal = cred['principalAmount']
    term = cred['termInMonths']
    freq = cred['paymentFrequency']
    disb_date = cred['disbursementDate']
    
    # Calcular proyección
    interest_rate = 0.10 # 10% mensual
    total_interest = principal * interest_rate * term
    total_amount = principal + total_interest
    
    if freq == 'weekly':
        num_installments = term * 4
        days_to_first = 7
    elif freq == 'biweekly':
        num_installments = term * 2
        days_to_first = 15
    else:
        num_installments = term
        days_to_first = 30
        
    first_inst_date = add_commercial_days(disb_date, days_to_first)
    
    inst_amount = math.ceil(total_amount / num_installments)
    principal_portion = math.ceil(principal / num_installments)
    interest_portion = inst_amount - principal_portion
    
    due_dates = generate_installment_dates(first_inst_date, num_installments, freq)
    
    installments = []
    accumulated_principal = 0.0
    accumulated_interest = 0.0
    
    for i in range(num_installments):
        is_last = (i == num_installments - 1)
        if is_last:
            curr_principal = principal - accumulated_principal
            curr_interest = total_interest - accumulated_interest
            curr_scheduled = curr_principal + curr_interest
        else:
            curr_principal = principal_portion
            curr_interest = interest_portion
            curr_scheduled = inst_amount
            
            accumulated_principal += curr_principal
            accumulated_interest += curr_interest
            
        installments.append({
            'installmentNumber': i + 1,
            'dueDate': due_dates[i],
            'principalPortion': curr_principal,
            'interestPortion': curr_interest,
            'scheduledAmount': curr_scheduled,
            'status': 'pending',
            'paidAmount': 0.0,
            'remainingAmount': curr_scheduled,
            'isMoraActive': False,
            'moraStartDate': due_dates[i],
            'dailyMoraRate': 0.007,
            'moraBase': curr_scheduled,
            'accumulatedMora': 0.0,
            'moraPaid': 0.0,
            'isConsolidated': False,
            'consolidatedIntoInstallment': None,
            'consolidationDate': None,
            'createdAt': datetime.datetime.now(),
            'updatedAt': disb_date,
            'alarmHour': 9,
            'alarmMinute': 0,
            'isAlarmEnabled': True,
            'isAlarmSilent': False
        })
        
    # Obtener pagos para este crédito
    cred_payments = [p for p in all_payments if p['creditId'] == c_id]
    
    # Ordenar por fecha_2
    cred_payments.sort(key=lambda x: x['fecha_2'])
    
    # Procesar pagos en cascada
    total_paid = 0.0
    total_paid_principal = 0.0
    total_paid_interest = 0.0
    
    payment_docs = []
    
    for pay in cred_payments:
        amt = pay['valor']
        cuota_no = pay['cuota_num']
        
        # Buscar fecha de vencimiento de la cuota objetivo
        pay_date = pay['fecha_2']
        target_inst = None
        for inst in installments:
            if inst['installmentNumber'] == cuota_no:
                target_inst = inst
                break
        if target_inst is not None:
            pay_date = target_inst['dueDate'] # Forzar para evitar moras
            
        remaining_pay = amt
        applied_int = 0.0
        applied_pri = 0.0
        affected_nums = []
        breakdowns = {}
        
        # Cascada de amortización
        for inst in installments:
            if inst['status'] == 'paid' or inst['remainingAmount'] <= 0:
                continue
            if remaining_pay <= 0:
                break
                
            sched_amt = inst['scheduledAmount']
            cap_ratio = inst['principalPortion'] / sched_amt if sched_amt > 0 else 0.0
            int_ratio = inst['interestPortion'] / sched_amt if sched_amt > 0 else 0.0
            
            int_paid_so_far = inst['paidAmount'] * int_ratio
            pri_paid_so_far = inst['paidAmount'] * cap_ratio
            
            unpaid_int = max(0.0, inst['interestPortion'] - int_paid_so_far)
            unpaid_pri = max(0.0, inst['principalPortion'] - pri_paid_so_far)
            total_unpaid = unpaid_int + unpaid_pri
            
            inst_applied_int = 0.0
            inst_applied_pri = 0.0
            
            if total_unpaid > 0 and remaining_pay > 0:
                if remaining_pay >= total_unpaid:
                    inst_applied_int = unpaid_int
                    inst_applied_pri = unpaid_pri
                    remaining_pay -= total_unpaid
                else:
                    target_p = remaining_pay * cap_ratio
                    target_i = remaining_pay * int_ratio
                    
                    if target_p > unpaid_pri:
                        inst_applied_pri = unpaid_pri
                        inst_applied_int = remaining_pay - unpaid_pri
                    elif target_i > unpaid_int:
                        inst_applied_int = unpaid_int
                        inst_applied_pri = remaining_pay - unpaid_int
                    else:
                        inst_applied_pri = target_p
                        inst_applied_int = target_i
                    remaining_pay = 0.0
                    
            if inst_applied_int > 0 or inst_applied_pri > 0:
                applied_int += inst_applied_int
                applied_pri += inst_applied_pri
                affected_nums.append(inst['installmentNumber'])
                breakdowns[str(inst['installmentNumber'])] = {
                    'mora': 0.0,
                    'interes': inst_applied_int,
                    'capital': inst_applied_pri
                }
                
                new_paid_amt = inst['paidAmount'] + inst_applied_int + inst_applied_pri
                new_rem = inst['scheduledAmount'] - new_paid_amt
                
                inst['paidAmount'] = new_paid_amt
                inst['remainingAmount'] = new_rem
                inst['moraBase'] = new_rem
                inst['status'] = 'paid' if new_rem <= 0 else 'partial'
                inst['updatedAt'] = pay_date
                
        actual_applied = amt - remaining_pay
        total_paid += actual_applied
        total_paid_principal += applied_pri
        total_paid_interest += applied_int
        
        # Payment doc dict
        pay_doc = {
            'registeredByUid': user_uid,
            'paymentDate': pay_date,
            'amountReceived': amt,
            'appliedToMora': 0.0,
            'appliedToInterest': applied_int,
            'appliedToPrincipal': applied_pri,
            'affectedInstallmentNumbers': affected_nums,
            'installmentBreakdowns': breakdowns,
            'paymentMethod': 'cash',
            'receiptNumber': None,
            'notes': "Migración de Google Sheets",
            'createdAt': datetime.datetime.now()
        }
        payment_docs.append(pay_doc)
        
    outstanding = max(0.0, total_amount - total_paid_principal - total_paid_interest)
    paid_count = sum(1 for inst in installments if inst['status'] == 'paid')
    
    cred_status = 'completed' if outstanding <= 0 else 'active'
    
    # Modelo del Crédito final
    credit_data = {
        'clientId': cl_id,
        'createdByUid': user_uid,
        'principalAmount': principal,
        'monthlyInterestRate': interest_rate,
        'dailyMoraRate': 0.007,
        'termInMonths': term,
        'paymentFrequency': freq,
        'totalInterest': total_interest,
        'totalAmount': total_amount,
        'installmentAmount': inst_amount,
        'numberOfInstallments': num_installments,
        'disbursementDate': disb_date,
        'firstInstallmentDate': first_inst_date,
        'status': cred_status,
        'paidInstallments': paid_count,
        'totalPaid': total_paid,
        'totalPaidPrincipal': total_paid_principal,
        'totalPaidInterest': total_paid_interest,
        'totalPaidMora': 0.0,
        'outstandingBalance': outstanding,
        'currentConsolidatedDebt': 0.0,
        'accumulatedMora': 0.0,
        'notes': f"Crédito migrado desde Excel (Código original: {c_id})",
        'createdAt': datetime.datetime.now(),
        'updatedAt': datetime.datetime.now()
    }
    
    # Guardar en Firestore de forma atómica en un único batch por crédito
    print(f"  Guardando crédito {c_id} de {c_name} (Pagos aplicados: {len(payment_docs)}, Balance pendiente: ${outstanding:,.0f})...")
    
    batch = db.batch()
    credit_ref = db.collection("credits").document(c_id)
    batch.set(credit_ref, credit_data)
    
    # Guardar cuotas
    for inst in installments:
        inst_ref = credit_ref.collection("installments").document()
        inst_data = inst.copy()
        inst_data['installmentId'] = inst_ref.id
        batch.set(inst_ref, inst_data)
        
    # Guardar pagos
    for p_doc in payment_docs:
        pay_ref = credit_ref.collection("payments").document()
        p_data = p_doc.copy()
        p_data['paymentId'] = pay_ref.id
        batch.set(pay_ref, p_data)
        
    batch.commit()

print("\n==================================================")
print("[OK] MIGRACION FINALIZADA CON EXITO")
print("==================================================")
