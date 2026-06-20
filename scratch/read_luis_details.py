import openpyxl

excel_path = r"C:\Users\Windows 10\Desktop\projects Antigravity\Copia de BANX.xlsx"
print(f"Loading Excel file: {excel_path}")
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Search 'Credito' sheet
print("\n=== SEARCHING IN SHEET 'Credito' ===")
ws_credito = wb['Credito']
headers_cred = [cell.value for cell in ws_credito[1]]
print("Headers:", headers_cred)
for r_idx in range(2, ws_credito.max_row + 1):
    row_vals = [ws_credito.cell(row=r_idx, column=c_idx).value for c_idx in range(1, ws_credito.max_column + 1)]
    code = str(row_vals[0]).strip() if row_vals[0] is not None else ""
    cedula = str(row_vals[4]).strip() if row_vals[4] is not None else ""
    name = str(row_vals[3]).strip() if row_vals[3] is not None else ""
    if "26021001" in [code, cedula] or "luis carlos" in name.lower() or "montenegro" in name.lower():
        print(f"Row {r_idx}: ID={code}, Name={name}, Cedula={cedula}, Value={row_vals[5]}, Term={row_vals[6]}, Date={row_vals[1]}, Freq={row_vals[8]}")

# 2. Search 'Clientes' sheet
print("\n=== SEARCHING IN SHEET 'Clientes' ===")
ws_clientes = wb['Clientes']
headers_cli = [cell.value for cell in ws_clientes[1]]
print("Headers:", headers_cli)
for r_idx in range(2, ws_clientes.max_row + 1):
    row_vals = [ws_clientes.cell(row=r_idx, column=c_idx).value for c_idx in range(1, ws_clientes.max_column + 1)]
    cedula = str(row_vals[0]).strip() if row_vals[0] is not None else ""
    name = str(row_vals[1]).strip() if row_vals[1] is not None else ""
    if "26021001" in [cedula] or "luis carlos" in name.lower() or "montenegro" in name.lower():
        print(f"Row {r_idx}: {row_vals}")

# 3. Search 'Pago Cuota' sheet
print("\n=== SEARCHING IN SHEET 'Pago Cuota' ===")
ws_pago = wb['Pago Cuota']
headers_pago = [cell.value for cell in ws_pago[1]]
print("Headers:", headers_pago)
for r_idx in range(2, ws_pago.max_row + 1):
    row_vals = [ws_pago.cell(row=r_idx, column=c_idx).value for c_idx in range(1, ws_pago.max_column + 1)]
    code = str(row_vals[2]).strip() if row_vals[2] is not None else ""
    name = str(row_vals[3]).strip() if row_vals[3] is not None else ""
    if code == "26021001" or "luis carlos" in name.lower() or "montenegro" in name.lower():
        print(f"Row {r_idx}: SchedDate={row_vals[0]}, PayDate={row_vals[1]}, Code={code}, Name={name}, Amount={row_vals[4]}, Cuota={row_vals[8]}, Other={row_vals[5:8]}")
