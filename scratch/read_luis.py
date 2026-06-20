import openpyxl

excel_path = r"C:\Users\Windows 10\Desktop\projects Antigravity\Copia de BANX.xlsx"
print(f"Loading {excel_path}...")
wb = openpyxl.load_workbook(excel_path, data_only=True)

for sheet_name in wb.sheetnames:
    print(f"\n--- Checking Sheet: {sheet_name} ---")
    ws = wb[sheet_name]
    header = [cell.value for cell in ws[1]]
    print(f"Header: {header}")
    
    rows_found = 0
    for r_idx in range(2, ws.max_row + 1):
        row_vals = [ws.cell(row=r_idx, column=c_idx).value for c_idx in range(1, ws.max_column + 1)]
        # Check if 26021001 or "luis carlos" is in the row values
        row_str = " ".join([str(v) for v in row_vals if v is not None]).lower()
        if "26021001" in row_str or "luis carlos" in row_str or "montenegro" in row_str:
            print(f"Row {r_idx}: {row_vals}")
            rows_found += 1
    print(f"Found {rows_found} rows in {sheet_name}")
